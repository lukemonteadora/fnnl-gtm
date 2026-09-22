#!/usr/bin/env python3
"""
social_proof_bank_to_xlsx.py

Generate the client-facing xlsx mirror of a client's Social Proof + Trigger Bank
from the markdown source of truth (profiles/<slug>/social-proof-bank.md).

The markdown is the source of truth. This produces a two-sheet workbook
(Case Studies, Trigger Bank) matching the example layout so it can be shared
with the client. Regenerate after any markdown update; never hand-edit the xlsx.

Usage:
    python social_proof_bank_to_xlsx.py --client <business-slug>
    python social_proof_bank_to_xlsx.py --client <slug> --in <path.md> --out <path.xlsx>

Dependency: openpyxl  (pip install openpyxl)

Parsing model: the markdown holds two GitHub-flavored pipe tables. Each table is
assigned to a sheet by the nearest preceding heading:
  - a heading containing "Case Stud"  -> Case Studies sheet (Bank 1)
  - a heading containing "Trigger"    -> Trigger Bank sheet (Bank 2)
Column headers are taken verbatim from each table's header row, humanized for display.
"""

import argparse
import os
import re
import sys


def repo_paths(script_path):
    scripts_dir = os.path.dirname(os.path.abspath(script_path))
    coldoutbound_root = os.path.dirname(scripts_dir)  # scripts/ -> outbound/
    profiles_dir = os.path.join(coldoutbound_root, "profiles")
    return coldoutbound_root, profiles_dir


def humanize(header):
    """snake_case / lower header -> Display Case, keeping the example's wording."""
    special = {
        "icp_segment": "ICP Segment",
        "service_line": "Service Line",
        "one_liner": "Case Study One-Liner (for copy)",
        "poke_the_bear": "Poke the Bear Question",
        "named_client": "Named Client",
        "naming_ok": "Name OK?",
        "proof_status": "Proof Status",
        "relevant_icps": "Relevant ICP(s)",
        "why_it_works": "Why It Works",
        "trigger_signal": "Trigger Signal",
        "canonical_column": "Canonical Column",
        "data_backed": "Data-Backed?",
    }
    if header in special:
        return special[header]
    return " ".join(w.capitalize() for w in re.split(r"[_\s]+", header.strip()))


def split_row(line):
    """Split a markdown table row into trimmed cells, dropping the outer pipes."""
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def is_separator(line):
    """True for a markdown header separator row like | --- | :--: |."""
    cells = split_row(line)
    return len(cells) > 0 and all(re.fullmatch(r":?-{2,}:?", c) is not None for c in cells if c != "")


def parse_tables(md_text):
    """
    Return {'case_studies': (headers, rows), 'trigger_bank': (headers, rows)}.
    A table is a run of consecutive lines starting with '|'. It is bound to the
    most recent heading; classification is by 'Case Stud' vs 'Trigger'.
    """
    lines = md_text.splitlines()
    current_bucket = None
    tables = {"case_studies": None, "trigger_bank": None}
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if stripped.startswith("#"):
            low = stripped.lower()
            if "case stud" in low:
                current_bucket = "case_studies"
            elif "trigger" in low:
                current_bucket = "trigger_bank"
            i += 1
            continue
        if stripped.startswith("|"):
            # collect the whole table block
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            if len(block) >= 2 and is_separator(block[1]):
                headers = split_row(block[0])
                rows = [split_row(r) for r in block[2:] if r.strip()]
                bucket = current_bucket
                # fall back: infer by header names if no heading told us
                if bucket is None:
                    joined = " ".join(headers).lower()
                    if "service_line" in joined or "one_liner" in joined:
                        bucket = "case_studies"
                    elif "trigger_signal" in joined or "data_backed" in joined:
                        bucket = "trigger_bank"
                if bucket:
                    if tables[bucket] is None:
                        tables[bucket] = (headers, rows)
                    else:
                        # Same bank, another SEGMENT's table (the bank is
                        # segment-aware: one Case Studies / Trigger table per
                        # segment section). Keep the first table's headers and
                        # append this segment's rows so no segment is dropped.
                        prev_headers, prev_rows = tables[bucket]
                        tables[bucket] = (prev_headers, prev_rows + rows)
            continue
        i += 1
    return tables


def build_workbook(client, tables, out_path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = Workbook()

    header_fill = PatternFill("solid", fgColor="1F3B57")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    title_font = Font(bold=True, size=13, color="1F3B57")
    wrap = Alignment(wrap_text=True, vertical="top")
    thin = Side(style="thin", color="D0D0D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    sheet_specs = [
        ("Case Studies", "case_studies",
         f"{client} - Case Studies (one row per ICP per service line)"),
        ("Trigger Bank", "trigger_bank",
         f"{client} - Trigger Bank (add rows as you discover new signals)"),
    ]

    first = True
    written_any = False
    for sheet_name, key, title in sheet_specs:
        table = tables.get(key)
        if not table:
            continue
        headers, rows = table
        ws = wb.active if first else wb.create_sheet()
        if first:
            ws.title = sheet_name
            first = False
        else:
            ws.title = sheet_name
        written_any = True

        ncol = len(headers)
        # Title row
        ws.cell(row=1, column=1, value=title).font = title_font
        if ncol > 1:
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)

        # Header row (row 2)
        for c, h in enumerate(headers, start=1):
            cell = ws.cell(row=2, column=c, value=humanize(h))
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left")
            cell.border = border

        # Data rows
        for r, row in enumerate(rows, start=3):
            for c in range(ncol):
                val = row[c] if c < len(row) else ""
                cell = ws.cell(row=r, column=c + 1, value=val)
                cell.alignment = wrap
                cell.border = border

        # Column widths (wider for copy columns)
        wide = {"one_liner", "poke_the_bear", "why_it_works", "result"}
        for c, h in enumerate(headers, start=1):
            letter = ws.cell(row=2, column=c).column_letter
            ws.column_dimensions[letter].width = 46 if h in wide else 20

        ws.freeze_panes = "A3"

    if not written_any:
        print("ERROR: no Case Studies or Trigger Bank table found in the markdown.",
              file=sys.stderr)
        return False

    wb.save(out_path)
    return True


# ---------------------------------------------------------------------------
# Client-fill template (--template): the blank Google-Sheet-shaped workbook we
# send the client when self-gather comes up short. Raw facts only, never copy.
# Industry-mixed example rows double as the fill-guide. Canonical structure is
# outbound/reference/social-proof-bank-client-fill-template.md.
# ---------------------------------------------------------------------------

CLIENT_FILL_HEADERS = [
    "Which customer?",
    "Can we name them?",
    "What did you do for them?",
    "Their industry / type",
    "Where",
    "The result (a real number)",
    "Over what period",
    "What it looked like before (optional)",
]

# Industries deliberately mixed so the sheet is obviously not vertical-specific.
CLIENT_FILL_EXAMPLES = [
    ["Northwind Software", "Yes", "Onboarding automation module", "B2B SaaS", "US",
     "Cut churn 22%", "2 quarters", "Manual onboarding, customers dropping in month 2"],
    ["Ardent Components", "Describe only", "Supply-planning rollout", "Manufacturing", "UK & DE",
     "Reduced lead time by 18 days", "6 months", "Orders slipping, planning on spreadsheets"],
    ["Riverside Clinic", "Yes", "No-show reduction program", "Healthcare provider", "US",
     "Cut no-show rate from 19% to 7%", "12 months", "1 in 5 appointments missed"],
    ["Beacon Home Services", "Yes", "Local demand + booking system", "Home services", "Texas",
     "Booked 40% more jobs", "90 days", "Leads going cold, no follow-up system"],
    ["Loom & Co", "Yes", "Repeat-purchase program", "Ecommerce / DTC", "US & CA",
     "Grew repeat-purchase revenue 3x", "1 year", "One-and-done buyers, no retention motion"],
    ["Halden Advisory", "No (NDA)", "Compliance review", "Professional services", "UK",
     "Recovered 480k in one review", "One engagement", "Recoverable costs sitting undetected"],
]

TIMING_QUESTIONS = [
    "When does a customer usually realize they need you? (the moment or event that makes them ready)",
    "Any seasonal or deadline timing that drives demand? (budget cycles, seasons, regulatory dates)",
    "Who are your main competitors? (names, so we can spot prospects using them)",
    "Any events, webinars, or conferences you'll be at soon?",
]

INSTRUCTIONS = [
    ("What this is",
     "A quick place to give us your real customer results so we can write outbound that "
     "names real clients and real numbers. We've already filled in everything we could find "
     "from your site and our calls. We just need you to confirm those and fill the blanks."),
    ("Why it matters",
     "Cold emails that name a real client and a real result get replies. Vague ones "
     "(\"we helped a leading firm\") get ignored. We'd rather use your real numbers than "
     "guess, so this is the highest-leverage 15 minutes you can give the campaign."),
    ("How to fill it",
     "Go to the \"Your results\" tab. Each row is one customer result. Fill the columns "
     "left to right. You only give us the facts. We write the actual emails. If you can't "
     "name a customer, say so in the \"Can we name them?\" column and we'll describe them "
     "instead. If a number is approximate, give your best real figure. Never invent one. "
     "The example rows below show the shape. Delete them or leave them, either is fine."),
]


def build_template_workbook(client, out_path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = Workbook()
    header_fill = PatternFill("solid", fgColor="1F3B57")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    title_font = Font(bold=True, size=13, color="1F3B57")
    label_font = Font(bold=True, size=11, color="1F3B57")
    example_fill = PatternFill("solid", fgColor="EEF3F8")
    wrap = Alignment(wrap_text=True, vertical="top")
    thin = Side(style="thin", color="D0D0D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ncol = len(CLIENT_FILL_HEADERS)

    # --- Tab 1: Start here (instructions + example rows) ---
    ws1 = wb.active
    ws1.title = "Start here"
    ws1.cell(1, 1, f"{client} - Proof + Results (for your outbound)").font = title_font
    ws1.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    r = 3
    for label, text in INSTRUCTIONS:
        ws1.cell(r, 1, label).font = label_font
        cell = ws1.cell(r + 1, 1, text)
        cell.alignment = wrap
        ws1.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=ncol)
        ws1.row_dimensions[r + 1].height = 48
        r += 3
    ws1.cell(r, 1, "Example rows (the shape to follow - delete or leave):").font = label_font
    r += 1
    for c, h in enumerate(CLIENT_FILL_HEADERS, start=1):
        cell = ws1.cell(r, c, h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = border
    for row in CLIENT_FILL_EXAMPLES:
        r += 1
        for c, val in enumerate(row, start=1):
            cell = ws1.cell(r, c, val)
            cell.alignment = wrap
            cell.border = border
            cell.fill = example_fill
    for c in range(1, ncol + 1):
        ws1.column_dimensions[ws1.cell(r, c).column_letter].width = 26

    # --- Tab 2: Your results (grid to fill) ---
    ws2 = wb.create_sheet("Your results")
    ws2.cell(1, 1, "Fill one row per customer result. Facts only - we write the emails. "
                   "Confirm any rows we pre-filled, and add your own below.").font = label_font
    ws2.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    ws2.row_dimensions[1].height = 30
    ws2.cell(1, 1).alignment = wrap
    for c, h in enumerate(CLIENT_FILL_HEADERS, start=1):
        cell = ws2.cell(2, c, h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = border
    for rr in range(3, 3 + 15):  # 15 blank formatted rows to fill
        for c in range(1, ncol + 1):
            ws2.cell(rr, c, "").border = border
    for c in range(1, ncol + 1):
        ws2.column_dimensions[ws2.cell(2, c).column_letter].width = 26
    ws2.freeze_panes = "A3"

    # --- Tab 3: Timing + competitors (optional) ---
    ws3 = wb.create_sheet("Timing + competitors")
    ws3.cell(1, 1, "Optional. Skip if the campaign only needs proof.").font = label_font
    ws3.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)
    rr = 3
    for q in TIMING_QUESTIONS:
        qc = ws3.cell(rr, 1, q)
        qc.alignment = wrap
        qc.font = label_font
        ans = ws3.cell(rr, 2, "")
        ans.border = border
        rr += 2
    ws3.column_dimensions["A"].width = 52
    ws3.column_dimensions["B"].width = 60

    wb.save(out_path)
    return True


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--client", help="business-slug (resolves the default in/out paths)")
    ap.add_argument("--in", dest="in_path", help="explicit path to social-proof-bank.md")
    ap.add_argument("--out", dest="out_path", help="explicit path to write the .xlsx")
    ap.add_argument("--template", action="store_true",
                    help="emit the blank CLIENT-FILL template (3 tabs + example rows) to send "
                         "the client, instead of mirroring a populated bank")
    args = ap.parse_args()

    _, profiles_dir = repo_paths(__file__)

    # --- Client-fill template mode: no populated bank required ---
    if args.template:
        out_path = args.out_path
        if not out_path:
            if not args.client:
                ap.error("provide --client, or --out explicitly, with --template")
            out_path = os.path.join(profiles_dir, args.client, "social-proof-bank-client-fill.xlsx")
        out_dir = os.path.dirname(out_path)
        if out_dir:
            os.makedirs(out_dir, exist_ok=True)
        build_template_workbook(args.client or "Client", out_path)
        print(f"Wrote client-fill template: {out_path}")
        print("  Tabs: Start here (with examples) | Your results | Timing + competitors")
        return

    in_path = args.in_path
    out_path = args.out_path
    if not in_path:
        if not args.client:
            ap.error("provide --client, or --in/--out explicitly")
        in_path = os.path.join(profiles_dir, args.client, "social-proof-bank.md")
    if not out_path:
        client = args.client or "client"
        base_dir = os.path.dirname(in_path)
        out_path = os.path.join(base_dir, "social-proof-bank.xlsx")

    if not os.path.isfile(in_path):
        print(f"ERROR: markdown not found: {in_path}", file=sys.stderr)
        sys.exit(2)

    with open(in_path, "r", encoding="utf-8") as f:
        md_text = f.read()

    tables = parse_tables(md_text)
    ok = build_workbook(args.client or "Client", tables, out_path)
    if not ok:
        sys.exit(1)

    cs = tables.get("case_studies")
    tb = tables.get("trigger_bank")
    print(f"Wrote {out_path}")
    print(f"  Case Studies rows: {len(cs[1]) if cs else 0}")
    print(f"  Trigger Bank rows: {len(tb[1]) if tb else 0}")


if __name__ == "__main__":
    main()
